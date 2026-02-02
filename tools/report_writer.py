import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Border, Side, Alignment
from tools.dedupe import remove_duplicates_df


def sort_csv_to_xlsx(fn: str) -> str:
    df = pd.read_csv(fn, header=None, parse_dates=[0])

    df = remove_duplicates_df(df)
    output = fn.replace(".csv", ".xlsx")
    df.to_excel(output, index=False, header=False)

    color_map = {
        "green": PatternFill(start_color="92D050", end_color="92D050", fill_type="solid"),
        "darkblue": PatternFill(start_color="9DC3E6", end_color="9DC3E6", fill_type="solid"),
        "orange": PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid"),
        "lightblue": PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid"),
        "grey": PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid"),
        "peach": PatternFill(start_color="FBE5D6", end_color="FBE5D6", fill_type="solid"),
        "yellow": PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid"),
    }

    color_keywords = [
        ("green", ["tab", "bud", "watch", "xr"]),
        ("darkblue", ["s21"]),
        ("orange", ["s22"]),
        ("lightblue", ["s23"]),
        ("grey", ["s24"]),
        ("peach", ["s25"]),
        ("yellow", ["flip", "fold"]),
    ]

    color_priority = {
        "green": 7,
        "darkblue": 6,
        "orange": 5,
        "lightblue": 4,
        "grey": 3,
        "peach": 2,
        "yellow": 1,
    }

    def get_row_color(row):
        last_color = None
        for color_name, keywords_list in color_keywords:
            for keyword in keywords_list:
                if any(keyword.lower() in str(row[col]).lower() for col in [4, 3, 2]):
                    last_color = color_name
        return last_color

    df["__color__"] = df.apply(get_row_color, axis=1)
    df["__priority__"] = df["__color__"].map(color_priority).fillna(999)
    df = df.sort_values("__priority__").reset_index(drop=True)
    df_to_save = df.drop(columns="__priority__")
    df_to_save.to_excel(output, index=False, header=False)

    wb = load_workbook(output)
    ws = wb.active

    thin = Side(border_style="thin", color="000000")
    all_border = Border(top=thin, left=thin, right=thin, bottom=thin)
    wrap_alignment = Alignment(wrap_text=True)

    for idx, row in df.iterrows():
        color_name = row["__color__"]
        ws.row_dimensions[idx + 1].height = 13.8

        for col in range(1, ws.max_column):
            cell = ws.cell(row=idx + 1, column=col)
            if color_name and col == 3:
                cell.fill = color_map[color_name]
            cell.border = all_border
            cell.alignment = wrap_alignment

    ws.delete_cols(ws.max_column)
    wb.save(output)
    return output
